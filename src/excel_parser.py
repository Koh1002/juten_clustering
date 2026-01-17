"""
Excel解析モジュール
- 結合セル解除
- データ抽出・成形
- カテゴリ列探索
"""

import io
import os
from pathlib import Path
from typing import List, Dict, Tuple, Optional, Union, BinaryIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, column_index_from_string

from .utils import (
    logger,
    normalize_text,
    is_summary_row,
    is_detail_row,
    extract_category,
    get_category_name,
    join_non_empty_cells,
    calculate_category_column_score
)


class ExcelParseError(Exception):
    """Excel解析エラー"""
    pass


class ExcelParser:
    """Excel解析クラス"""

    # 対象シート名
    TARGET_SHEET = "提出用"

    # カテゴリ列探索範囲（P列〜V列 = 16〜22）
    CATEGORY_SEARCH_START_COL = 16  # P列
    CATEGORY_SEARCH_END_COL = 22    # V列
    CATEGORY_SEARCH_ROWS = 200      # 探索行数

    def __init__(self):
        self.parse_results = []
        self.category_column_info = {}
        self.errors = []
        self.warnings = []

    def load_files(
        self,
        file_paths: Optional[List[str]] = None,
        file_objects: Optional[List[BinaryIO]] = None
    ) -> pd.DataFrame:
        """
        複数のExcelファイルを読み込んで結合
        file_paths: ファイルパスのリスト（ローカル実行用）
        file_objects: ファイルオブジェクトのリスト（アップロード用）
        """
        self.parse_results = []
        self.errors = []
        self.warnings = []
        all_records = []

        sources = []
        if file_paths:
            sources.extend([('path', p) for p in file_paths])
        if file_objects:
            sources.extend([('object', f) for f in file_objects])

        if not sources:
            raise ExcelParseError("ファイルが指定されていません")

        for source_type, source in sources:
            try:
                if source_type == 'path':
                    filename = os.path.basename(source)
                    wb = load_workbook(source, data_only=True)
                else:
                    filename = getattr(source, 'name', 'uploaded_file.xlsx')
                    # BytesIOに変換
                    if hasattr(source, 'read'):
                        content = source.read()
                        if hasattr(source, 'seek'):
                            source.seek(0)
                        wb = load_workbook(io.BytesIO(content), data_only=True)
                    else:
                        wb = load_workbook(source, data_only=True)

                # シート存在確認
                if self.TARGET_SHEET not in wb.sheetnames:
                    self.errors.append(f"「{filename}」に「{self.TARGET_SHEET}」シートがありません（スキップ）")
                    continue

                ws = wb[self.TARGET_SHEET]

                # 結合セル解除
                self._unmerge_and_fill(ws)

                # データ抽出
                records = self._extract_records(ws, filename)
                all_records.extend(records)

                self.parse_results.append({
                    'filename': filename,
                    'record_count': len(records),
                    'status': 'success'
                })

            except Exception as e:
                logger.error(f"Error loading {source}: {e}")
                self.errors.append(f"「{filename if 'filename' in dir() else source}」の読み込みエラー: {str(e)}")

        if not all_records:
            raise ExcelParseError(
                "有効なレコードが抽出できませんでした。\n"
                "考えられる原因:\n"
                "- 「提出用」シートがない\n"
                "- 番号行（1. や (1) など）の形式が異なる\n"
                "- データが7行目以降に存在しない"
            )

        df = pd.DataFrame(all_records)
        return df

    def _unmerge_and_fill(self, ws):
        """
        結合セルを解除し、左上セルの値を範囲内にフィル
        """
        # 結合範囲のリストをコピー（イテレーション中に変更するため）
        merged_ranges = list(ws.merged_cells.ranges)

        for merged_range in merged_ranges:
            # 左上セルの値を取得
            min_row, min_col = merged_range.min_row, merged_range.min_col
            top_left_value = ws.cell(row=min_row, column=min_col).value

            # 結合解除
            ws.unmerge_cells(str(merged_range))

            # 範囲内の全セルに値をフィル
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    ws.cell(row=row, column=col).value = top_left_value

    def _find_category_column(self, ws, manual_column: Optional[str] = None) -> Tuple[int, dict]:
        """
        カテゴリ列を探索
        manual_column: 手動指定された列（例: "R"）
        """
        if manual_column:
            col_idx = column_index_from_string(manual_column)
            values = []
            for row in range(7, 7 + self.CATEGORY_SEARCH_ROWS):
                cell = ws.cell(row=row, column=col_idx)
                values.append(cell.value)
            score_info = calculate_category_column_score(values)
            score_info['column'] = manual_column
            score_info['column_index'] = col_idx
            return col_idx, score_info

        # 自動探索
        best_col = None
        best_score = -1
        best_info = {}
        all_candidates = []

        for col_idx in range(self.CATEGORY_SEARCH_START_COL, self.CATEGORY_SEARCH_END_COL + 1):
            col_letter = get_column_letter(col_idx)
            values = []
            for row in range(7, 7 + self.CATEGORY_SEARCH_ROWS):
                cell = ws.cell(row=row, column=col_idx)
                values.append(cell.value)

            score_info = calculate_category_column_score(values)
            score_info['column'] = col_letter
            score_info['column_index'] = col_idx
            all_candidates.append(score_info)

            if score_info['score'] > best_score:
                best_score = score_info['score']
                best_col = col_idx
                best_info = score_info

        self.category_column_info = {
            'best': best_info,
            'candidates': sorted(all_candidates, key=lambda x: x['score'], reverse=True)
        }

        if best_score < 0.01:  # ほぼマッチしない
            self.warnings.append(
                f"カテゴリ列が検出できませんでした（最高スコア: {best_score:.2%}）。"
                f"手動で列を指定してください。"
            )

        return best_col, best_info

    def _extract_records(
        self,
        ws,
        filename: str,
        manual_category_column: Optional[str] = None,
        inherit_category: bool = True
    ) -> List[Dict]:
        """
        ワークシートからレコードを抽出
        """
        records = []

        # 固定セル取得（D4: 店舗名, N4: ブロック番号）
        store_name = normalize_text(ws.cell(row=4, column=4).value)  # D4
        block_number = normalize_text(ws.cell(row=4, column=14).value)  # N4

        # カテゴリ列探索
        category_col, cat_info = self._find_category_column(ws, manual_category_column)

        # 現在の概要情報を保持
        current_summary = ""
        current_category = None
        last_category = None

        # 7行目以降を走査
        max_row = min(ws.max_row, 7 + self.CATEGORY_SEARCH_ROWS)

        for row_idx in range(7, max_row + 1):
            # E列（5列目）で概要行判定
            e_value = ws.cell(row=row_idx, column=5).value

            # F列（6列目）で詳細行判定
            f_value = ws.cell(row=row_idx, column=6).value

            # カテゴリ取得
            if category_col:
                cat_value = ws.cell(row=row_idx, column=category_col).value
                cat = extract_category(cat_value)
                if cat:
                    last_category = cat
            else:
                cat = None

            if is_summary_row(e_value):
                # 概要行の場合
                # F〜O列（6〜15列）から本文取得
                text_cells = [ws.cell(row=row_idx, column=c).value for c in range(6, 16)]
                current_summary = join_non_empty_cells(text_cells)

                # カテゴリ設定
                if cat:
                    current_category = cat
                elif inherit_category and last_category:
                    current_category = last_category
                else:
                    current_category = None

                # この概要に詳細がつくか確認（次行以降をチェック）
                has_detail = False
                for check_row in range(row_idx + 1, min(row_idx + 50, max_row + 1)):
                    check_e = ws.cell(row=check_row, column=5).value
                    check_f = ws.cell(row=check_row, column=6).value
                    if is_summary_row(check_e):
                        break  # 次の概要行に到達
                    if is_detail_row(check_f):
                        has_detail = True
                        break

                # 詳細がない場合は概要のみでレコード作成
                if not has_detail and current_summary:
                    records.append({
                        '店舗名': store_name,
                        'ブロック番号': block_number,
                        '取り組み宣言（概要）': current_summary,
                        '取り組み宣言（詳細）': '',
                        'カテゴリ': current_category or '',
                        'カテゴリ名': get_category_name(current_category),
                        '_source_file': filename
                    })

            elif is_detail_row(f_value):
                # 詳細行の場合
                # G〜O列（7〜15列）から本文取得
                text_cells = [ws.cell(row=row_idx, column=c).value for c in range(7, 16)]
                detail_text = join_non_empty_cells(text_cells)

                # カテゴリ更新
                if cat:
                    current_category = cat
                elif inherit_category and last_category:
                    current_category = last_category

                if detail_text:
                    records.append({
                        '店舗名': store_name,
                        'ブロック番号': block_number,
                        '取り組み宣言（概要）': current_summary,
                        '取り組み宣言（詳細）': detail_text,
                        'カテゴリ': current_category or '',
                        'カテゴリ名': get_category_name(current_category),
                        '_source_file': filename
                    })

        return records

    def get_category_column_info(self) -> dict:
        """カテゴリ列探索結果を取得"""
        return self.category_column_info

    def get_parse_results(self) -> List[dict]:
        """解析結果サマリを取得"""
        return self.parse_results

    def get_errors(self) -> List[str]:
        """エラーリストを取得"""
        return self.errors

    def get_warnings(self) -> List[str]:
        """警告リストを取得"""
        return self.warnings


def select_files_with_dialog() -> Tuple[List[str], str]:
    """
    tkinterダイアログでファイル/フォルダを選択
    戻り値: (ファイルパスリスト, 選択モード)
    """
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)

        # ファイル選択ダイアログ
        file_paths = filedialog.askopenfilenames(
            title="Excelファイルを選択（複数可）",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )

        root.destroy()

        if file_paths:
            return list(file_paths), 'files'

        # フォルダ選択を試行
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)

        folder_path = filedialog.askdirectory(title="フォルダを選択")

        root.destroy()

        if folder_path:
            # フォルダ内のExcelファイルを取得
            folder = Path(folder_path)
            excel_files = list(folder.glob("*.xlsx")) + list(folder.glob("*.xls"))
            return [str(f) for f in excel_files], 'folder'

        return [], 'cancelled'

    except Exception as e:
        logger.error(f"Dialog error: {e}")
        return [], 'error'
