"""
サンプルExcelファイル生成スクリプト
テスト用の「提出用」シートを含むExcelを作成
"""

import os
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def create_sample_excel(output_path: str):
    """サンプルExcelファイルを作成"""
    wb = Workbook()

    # デフォルトシートを「提出用」にリネーム
    ws = wb.active
    ws.title = "提出用"

    # D4: 店舗名（結合セル）
    ws.merge_cells('D4:G4')
    ws['D4'] = '東京本店'

    # N4: ブロック番号（結合セル）
    ws.merge_cells('N4:O4')
    ws['N4'] = 'B001'

    # ヘッダー行（6行目）- 参考用
    headers = {
        'E': '番号',
        'F': '概要',
        'G': '詳細',
        'R': 'カテゴリ'
    }
    for col, header in headers.items():
        ws[f'{col}6'] = header

    # サンプルデータ
    sample_data = [
        # (行番号, E列番号記号, F列詳細記号, F:O本文, R列カテゴリ)
        (7, '1.', None, '顧客満足度の向上を目指した接客サービスの改善', 'A'),
        (8, None, '(1)', '挨拶の徹底と笑顔での対応を全スタッフに教育する', 'A'),
        (9, None, '(2)', 'お客様からのフィードバックを毎週集計し改善に活かす', 'A'),
        (10, '2.', None, '売上目標達成に向けた販売戦略の策定', 'B'),
        (11, None, '(1)', '季節商品の早期展開と効果的な陳列を実施する', 'B'),
        (12, None, '(2)', 'リピーター獲得のためのポイントカード活用促進', 'B'),
        (13, '3.', None, 'スタッフの能力開発とキャリア支援', 'C'),
        (14, None, '(1)', '月1回の研修会開催とOJTプログラムの充実', 'C'),
        (15, '4.', None, '業務効率化によるコスト削減', 'E'),
        (16, None, '(1)', '在庫管理システムの導入と発注業務の最適化', 'E'),
        (17, None, '(2)', '省エネ施策の推進と光熱費の削減目標設定', 'E'),
    ]

    for row_num, e_val, f_val, text, category in sample_data:
        if e_val:
            ws[f'E{row_num}'] = e_val
            # 概要行はF:O結合
            ws.merge_cells(f'F{row_num}:O{row_num}')
            ws[f'F{row_num}'] = text
        elif f_val:
            ws[f'F{row_num}'] = f_val
            # 詳細行はG:O結合
            ws.merge_cells(f'G{row_num}:O{row_num}')
            ws[f'G{row_num}'] = text

        # カテゴリ（R列）
        ws[f'R{row_num}'] = f'【{category}】'

    # 「作成例」シートも追加（これは無視されるべき）
    ws2 = wb.create_sheet("作成例")
    ws2['A1'] = "このシートは読み込まれません"

    # 保存
    wb.save(output_path)
    print(f"サンプルExcelを作成しました: {output_path}")


def create_multiple_samples(output_dir: str, count: int = 3):
    """複数のサンプルExcelファイルを作成"""
    os.makedirs(output_dir, exist_ok=True)

    stores = ['東京本店', '大阪支店', '名古屋支店', '福岡支店', '札幌支店']
    blocks = ['B001', 'B002', 'B003', 'B004', 'B005']

    for i in range(count):
        wb = Workbook()
        ws = wb.active
        ws.title = "提出用"

        # 店舗情報
        store = stores[i % len(stores)]
        block = blocks[i % len(blocks)]

        ws.merge_cells('D4:G4')
        ws['D4'] = store

        ws.merge_cells('N4:O4')
        ws['N4'] = block

        # サンプルデータ（店舗ごとに少し違うデータ）
        base_data = [
            (7, '1.', None, f'{store}における顧客満足度向上施策', 'A'),
            (8, None, '(1)', '接客マニュアルの刷新と定期トレーニング', 'A'),
            (9, None, '(2)', '顧客アンケートの実施と分析', 'A'),
            (10, '2.', None, '売上拡大に向けた取り組み', 'B'),
            (11, None, '(1)', '新商品の積極的なプロモーション', 'B'),
            (12, '3.', None, '人材育成プログラムの強化', 'C'),
            (13, None, '(1)', '階層別研修の実施', 'C'),
            (14, None, '(2)', 'メンター制度の導入', 'C'),
            (15, '4.', None, '目標達成のための実行力強化', 'D'),
            (16, None, '(1)', '週次進捗会議の開催', 'D'),
            (17, '5.', None, 'コスト管理の徹底', 'E'),
            (18, None, '(1)', '経費削減目標の設定と進捗管理', 'E'),
        ]

        for row_num, e_val, f_val, text, category in base_data:
            if e_val:
                ws[f'E{row_num}'] = e_val
                ws.merge_cells(f'F{row_num}:O{row_num}')
                ws[f'F{row_num}'] = text
            elif f_val:
                ws[f'F{row_num}'] = f_val
                ws.merge_cells(f'G{row_num}:O{row_num}')
                ws[f'G{row_num}'] = text

            ws[f'R{row_num}'] = f'【{category}】'

        # 作成例シート
        ws2 = wb.create_sheet("作成例")
        ws2['A1'] = "サンプル"

        output_path = os.path.join(output_dir, f'sample_{store}.xlsx')
        wb.save(output_path)
        print(f"作成: {output_path}")


if __name__ == "__main__":
    import sys

    # 出力ディレクトリ
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    data_dir = os.path.join(project_dir, 'data')

    if len(sys.argv) > 1 and sys.argv[1] == '--multiple':
        create_multiple_samples(data_dir, 3)
    else:
        output_path = os.path.join(data_dir, 'sample_template.xlsx')
        os.makedirs(data_dir, exist_ok=True)
        create_sample_excel(output_path)
